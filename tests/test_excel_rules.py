"""Prove that workflows.xlsx really controls behaviour (not just routing)."""
import shutil
from pathlib import Path
import openpyxl
import pytest
from src.main import build_agent
from src.workflow.excel_parser import load_business_excel
from src.engine.rules import parse_rules, parse_steps, parse_tools

ROOT = Path(__file__).resolve().parents[1]
DATA = load_business_excel(ROOT / 'data' / 'sample_daily_business_data.xlsx')


def _edited_workbook(tmp_path, wid, column, fn):
    path = tmp_path / 'workflows.xlsx'
    shutil.copy(ROOT / 'data' / 'workflows.xlsx', path)
    wb = openpyxl.load_workbook(path)
    ws = wb['Workflows']
    headers = [c.value for c in ws[1]]
    col = headers.index(column) + 1
    for row in ws.iter_rows(min_row=2):
        if row[0].value == wid:
            row[col - 1].value = fn(row[col - 1].value)
    wb.save(path)
    return path


def test_parsers():
    assert parse_rules('Flag when price difference exceeds 10%') == {'percent_threshold': 10.0}
    assert parse_rules('time above 2 minutes')['time_threshold_seconds'] == 120
    assert parse_steps('A → B → C') == ['A', 'B', 'C']
    assert parse_tools('CSV reader; calculator') == ['CSV reader', 'calculator']


def test_changing_threshold_in_excel_changes_price_exceptions(tmp_path):
    base = build_agent()('Find products where vendor price differs by more than 10%.', DATA)[1]
    assert base['Exception'].sum() == 3
    path = _edited_workbook(tmp_path, 'WF002', 'Decision_Logic', lambda v: v.replace('10%', '15%'))
    result = build_agent(path)('Find products where vendor price differs by more than 15%.', DATA)[1]
    assert result['Exception'].sum() == 1          # only the 25% difference


def test_failure_threshold_in_excel_changes_performance_flags(tmp_path):
    base = build_agent()('Which workflows are failing most often?', DATA)[1]
    assert base['metrics']['Flagged'].any()
    path = _edited_workbook(tmp_path, 'WF010', 'Decision_Logic', lambda v: v.replace('10%', '99%'))
    result = build_agent(path)('Which workflows are failing most often?', DATA)[1]
    assert result['metrics']['Flagged'].sum() < base['metrics']['Flagged'].sum()


def test_unregistered_tool_in_excel_is_rejected(tmp_path):
    path = _edited_workbook(tmp_path, 'WF001', 'Tools_Required', lambda v: v + '; fax machine')
    with pytest.raises(ValueError, match='not registered'):
        build_agent(path)('Which products need restocking?', DATA)


def test_trace_contains_steps_defined_in_excel(tmp_path):
    path = _edited_workbook(tmp_path, 'WF001', 'Steps', lambda v: v + ' → send alert email')
    _, _, ctx = build_agent(path)('Which products need restocking?', DATA)
    assert any('send alert email' in t['step'] for t in ctx.trace)
